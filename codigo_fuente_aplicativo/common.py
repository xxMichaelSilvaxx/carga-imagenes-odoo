"""Shared helpers: data folders, progress/ETA, references and the worker-to-GUI bridge."""
import logging
import re
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path

APP_NAME = "Carga de imágenes a Odoo"
IMAGE_EXTENSIONS = (".webp", ".png", ".jpg", ".jpeg", ".gif", ".bmp")
SOURCE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".ppm", ".ico")
DEFAULT_SUFFIXES = "PROMO"
STATUS_KIND = {
    "Nueva": "ok", "Actualizada": "ok", "Convertida": "ok", "Se agregaría": "ok", "Se actualizaría": "ok",
    "Error": "error", "No existe en Odoo": "error",
    "Sin imagen en la carpeta": "warn", "Omitida": "warn", "Cancelada": "warn", "Copiada sin convertir": "warn",
}


def data_dir(*parts):
    base = Path.home() / "Documents"
    folder = (base if base.is_dir() else Path.home()) / "CargaImagenesOdoo"
    folder = folder.joinpath(*parts)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def resource_path(name):
    """File shipped with the app: next to the sources, or in the PyInstaller temp folder inside the .exe."""
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / name


def set_window_icon(window):
    import tkinter as tk

    try:
        if sys.platform == "win32":
            window.iconbitmap(default=str(resource_path("icon.ico")))
        else:
            window.icon_image = tk.PhotoImage(master=window, file=str(resource_path("icon.png")))
            window.iconphoto(True, window.icon_image)
    except tk.TclError:
        pass


def setup_logger():
    logger = logging.getLogger("carga_imagenes")
    if not logger.handlers:
        log_file = data_dir("logs") / f"carga_{datetime.now():%Y%m%d}.log"
        handler = logging.FileHandler(log_file, encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
    return logger


def fold(text):
    stripped = unicodedata.normalize("NFKD", str(text or "").replace("﻿", "").lower())
    return "".join(c for c in stripped if not unicodedata.combining(c)).strip()


def clean(value):
    return str(value or "").replace("﻿", "").replace("\xa0", " ").strip().strip('"').strip("'").strip()


def parse_suffixes(text):
    return tuple(part for part in (clean(item) for item in re.split(r"[,;\n]", text or "")) if part)


class Matcher:
    """Compare references ignoring case and accents; base() drops suffixes like _PROMOCIONALES."""

    def __init__(self, suffixes=()):
        words = "|".join(re.escape(fold(word).strip("_- .")) for word in suffixes if fold(word).strip("_- ."))
        self.pattern = re.compile(rf"[\s_.\-]*(?:{words})$") if words else None

    def exact(self, text):
        return fold(text)

    def base(self, text):
        key = fold(text)
        if not self.pattern:
            return key
        return self.pattern.sub("", key) or key


def read_references(path):
    """One internal reference per line; blanks, # comments and extra columns are ignored."""
    raw = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            raw = Path(path).read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    references, seen = [], set()
    for line in (raw or "").splitlines():
        line = clean(line)
        if not line or line.startswith("#"):
            continue
        line = clean(re.split(r"[;,\t]", line)[0])
        if line and fold(line) not in seen:
            seen.add(fold(line))
            references.append(line)
    return references


def human_time(seconds):
    seconds = max(int(seconds), 0)
    if seconds < 60:
        return f"{seconds} s"
    if seconds < 3600:
        return f"{seconds // 60} min {seconds % 60:02d} s"
    return f"{seconds // 3600} h {(seconds % 3600) // 60:02d} min"


class Progress:
    def __init__(self, total):
        self.total = total
        self.done = 0
        self.started = time.monotonic()

    def step(self, count=1):
        self.done += count

    def snapshot(self):
        elapsed = time.monotonic() - self.started
        rate = self.done / elapsed if elapsed > 0 else 0
        remaining = (self.total - self.done) / rate if rate else None
        return {
            "done": self.done,
            "total": self.total,
            "percent": 100 * self.done / self.total if self.total else 100,
            "rate": rate,
            "elapsed": human_time(elapsed),
            "eta": human_time(remaining) if remaining is not None else "calculando...",
        }


class Reporter:
    """Bridge from worker threads to the GUI queue, mirrored to the daily log file."""

    def __init__(self, events, logger):
        self.events = events
        self.logger = logger
        self.rows = []
        self.columns = ()

    def begin(self, columns):
        self.columns = columns
        self.rows = []
        self.events.put(("begin", columns))

    def info(self, text):
        self._log("info", text)

    def warning(self, text):
        self._log("warning", text)

    def error(self, text):
        self._log("error", text)

    def success(self, text):
        self._log("success", text)

    def _log(self, level, text):
        self.logger.log(logging.ERROR if level == "error" else logging.WARNING if level == "warning" else logging.INFO, text)
        self.events.put(("log", level, text))

    def status(self, text):
        self.events.put(("status", text))

    def progress(self, progress, extra=""):
        self.events.put(("progress", progress.snapshot(), extra))

    def row(self, *values):
        self.rows.append(values)
        self.logger.info(" | ".join(str(value) for value in values))
        self.events.put(("row", values))

    def write_report(self, kind, pending=()):
        stamp = f"{datetime.now():%Y%m%d_%H%M%S}"
        folder = data_dir("reportes")
        report = folder / f"{kind}_{stamp}.xlsx"
        write_workbook(report, self.columns, self.rows)
        if pending:
            pending_file = folder / f"{kind}_pendientes_{stamp}.txt"
            pending_file.write_text("\n".join(pending), encoding="utf-8")
            self.info(f"Referencias pendientes (para reintentar): {pending_file}")
        self.info(f"Reporte completo: {report}")
        return report

    def finished(self, summary, report=None):
        self.logger.info(summary)
        self.events.put(("done", summary, str(report) if report else ""))


def write_workbook(path, columns, rows):
    """Excel report where every cell is text, so codes like 0201 or 10E5 stay untouched."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {kind: PatternFill("solid", fgColor=color)
             for kind, color in (("ok", "E7F6EC"), ("error", "FDE8E8"), ("warn", "FFF4E0"))}
    status_index = columns.index("Estado") if "Estado" in columns else None
    book = Workbook()
    sheet = book.active
    sheet.title = "Resultados"
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="223047")
        cell.alignment = Alignment(vertical="center")
    widths = [len(column) for column in columns]
    for values in rows:
        sheet.append([str(value) for value in values])
        fill = fills.get(STATUS_KIND.get(values[status_index])) if status_index is not None else None
        for index, cell in enumerate(sheet[sheet.max_row]):
            cell.number_format = "@"
            if fill and index == status_index:
                cell.fill = fill
        widths = [max(width, len(str(value))) for width, value in zip(widths, values)]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 70)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = book.create_sheet("Resumen")
    summary.append(["Estado", "Cantidad"])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="223047")
    counts = {}
    for values in rows:
        if status_index is not None:
            counts[values[status_index]] = counts.get(values[status_index], 0) + 1
    for status, count in sorted(counts.items(), key=lambda item: -item[1]):
        summary.append([status, count])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 12
    book.save(path)
